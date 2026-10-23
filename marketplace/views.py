import json
import logging
import re
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from functools import wraps
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from django import forms
from django.core.exceptions import ValidationError
from django.core.validators import URLValidator
from django.http import HttpResponseForbidden, JsonResponse
from django.shortcuts import render, get_object_or_404, redirect
from django.conf import settings
from django.http import FileResponse, Http404, HttpResponse
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.admin.views.decorators import staff_member_required
from django.db import transaction
from django.db.models import Count, F, Q, Sum
from django.utils import timezone
from django.views.decorators.http import require_POST
from django.views.decorators.csrf import ensure_csrf_cookie
from pywebpush import WebPushException, webpush
from .models import (
    Apartment,
    Customer,
    FeaturedCampaign,
    Order,
    OrderItem,
    Product,
    ProductGroup,
    PushSubscription,
    StoreConfig,
    WhatsAppNotificationLog,
)

logger = logging.getLogger(__name__)


# ---------- Cart helpers ----------

def _get_cart(request):
    """Return session cart. Format: { 'product_id_str': quantity_int }"""
    return request.session.get('cart', {})


def _save_cart(request, cart):
    request.session['cart'] = cart
    request.session.modified = True


def _cart_total_items(cart):
    return sum(int(q) for q in cart.values())


# ---------- Views ----------
@ensure_csrf_cookie
def home(request):
    query = request.GET.get('q', '').strip()
    groups = ProductGroup.objects.all().order_by('name')
    search_results = []
    now = timezone.now()
    active_campaigns = FeaturedCampaign.objects.filter(
        is_active=True,
        starts_at__lte=now,
        ends_at__gte=now,
    )
    featured_products = Product.objects.filter(
        featured_in__in=active_campaigns,
        is_active=True,
        selling_price__gt=0,
    ).distinct()[:12]
    current_customer = None

    customer_id = request.session.get('customer_id')
    if customer_id:
        try:
            current_customer = Customer.objects.get(id=customer_id)
        except (Customer.DoesNotExist, ValueError, TypeError):
            request.session.pop('customer_id', None)

    frequent_products = []
    if current_customer:
        frequent_products = Product.objects.filter(
            orderitem__order__customer=current_customer,
            is_active=True,
            selling_price__gt=0,
        ).annotate(
            times_ordered=Count('orderitem')
        ).order_by('-times_ordered')[:6]

    if query:
        search_results = Product.objects.filter(
            Q(name__icontains=query) | Q(code__icontains=query),
            is_active=True,
            selling_price__gt=0,
        ).order_by('name')[:50]

    config = _get_store_config()
    return render(request, 'marketplace/home.html', {
        'groups': groups,
        'query': query,
        'search_results': search_results,
        'current_customer': current_customer,
        'frequent_products': frequent_products,
        'featured_products': featured_products,
        'config': config,
    })


def product_list(request, group_id):
    group = get_object_or_404(ProductGroup, id=group_id)
    products = Product.objects.filter(
        group=group, is_active=True, selling_price__gt=0
    ).order_by('name')
    return render(request, 'marketplace/product_list.html', {
        'group': group,
        'products': products,
        'config': _get_store_config(),
    })


def cart_count(request):
    """Returns just the total quantity in cart (as plain text)."""
    cart = _get_cart(request)
    return HttpResponse(str(_cart_total_items(cart)))


@require_POST
def add_to_cart(request, product_id):
    """Add one unit of the product to the cart. Responds with new count + HTMX trigger."""
    if not _get_store_config().is_open:
        return HttpResponseForbidden("Store is closed")

    product = get_object_or_404(Product, id=product_id, is_active=True)
    cart = _get_cart(request)
    key = str(product.id)
    cart[key] = int(cart.get(key, 0)) + 1
    _save_cart(request, cart)

    total = _cart_total_items(cart)
    response = HttpResponse(str(total))
    response['HX-Trigger'] = 'cartUpdated'  # Fires event that navbar listens for
    return response


# ---------- Cart page + updates ----------

def _get_store_config():
    config, _ = StoreConfig.objects.get_or_create(pk=1)
    return config


def _send_push_to_customer(customer, title, body, url):
    if not customer:
        return

    subscriptions = customer.push_subscriptions.all()
    if not subscriptions.exists():
        return
    if not settings.VAPID_PRIVATE_KEY or not settings.VAPID_ADMIN_EMAIL:
        logger.error('Push notifications are unavailable: VAPID settings are missing.')
        return

    payload = json.dumps({
        'title': title,
        'body': body,
        'url': url,
    })
    for subscription in subscriptions:
        try:
            webpush(
                subscription_info={
                    'endpoint': subscription.endpoint,
                    'keys': {
                        'p256dh': subscription.p256dh,
                        'auth': subscription.auth,
                    },
                },
                data=payload,
                vapid_private_key=settings.VAPID_PRIVATE_KEY,
                vapid_claims={
                    'sub': f'mailto:{settings.VAPID_ADMIN_EMAIL}',
                },
                ttl=3600,
                timeout=10,
            )
        except WebPushException as exc:
            status_code = (
                exc.response.status_code
                if exc.response is not None
                else None
            )
            if status_code in {404, 410}:
                subscription.delete()
                logger.info(
                    'Removed expired push subscription id=%s (HTTP %s).',
                    subscription.pk,
                    status_code,
                )
            else:
                logger.warning(
                    'Push delivery failed for subscription id=%s (HTTP %s).',
                    subscription.pk,
                    status_code or 'unknown',
                )


def store_open_required(view_func):
    @wraps(view_func)
    def wrapped_view(request, *args, **kwargs):
        config = _get_store_config()
        if not config.is_open:
            messages.error(request, config.closed_message)
            return redirect('cart_view')
        return view_func(request, *args, **kwargs)

    return wrapped_view


def _calculate_delivery_fee(subtotal, apartment=None, config=None):
    config = config or _get_store_config()
    if subtotal >= config.free_delivery_threshold:
        return Decimal('0.00')
    if apartment and apartment.delivery_fee_override is not None:
        return apartment.delivery_fee_override
    return config.delivery_fee


def _cart_details(request, apartment_id=None):
    """Build cart items and calculated order totals for the selected apartment."""
    cart = _get_cart(request)
    items = []
    subtotal = Decimal('0.00')
    for pid_str, qty in cart.items():
        try:
            product = Product.objects.get(id=int(pid_str), is_active=True)
        except Product.DoesNotExist:
            continue
        line_total = product.selling_price * int(qty)
        subtotal += line_total
        items.append({
            'product': product,
            'quantity': int(qty),
            'line_total': line_total,
        })
    config = _get_store_config()
    apartment = None
    if apartment_id:
        try:
            apartment = Apartment.objects.filter(pk=apartment_id).first()
        except (ValueError, TypeError):
            apartment = None
    delivery_fee = _calculate_delivery_fee(subtotal, apartment, config)
    grand_total = subtotal + delivery_fee
    amount_for_free_delivery = max(
        config.free_delivery_threshold - subtotal,
        Decimal('0.00'),
    ) if subtotal > 0 else Decimal('0.00')
    return (
        items,
        subtotal,
        delivery_fee,
        grand_total,
        amount_for_free_delivery,
        config,
    )


def cart_view(request):
    items, subtotal, delivery_fee, grand_total, amount_for_free_delivery, config = _cart_details(request)
    return render(request, 'marketplace/cart.html', {
        'items': items,
        'subtotal': subtotal,
        'delivery_fee': delivery_fee,
        'grand_total': grand_total,
        'amount_for_free_delivery': amount_for_free_delivery,
        'config': config,
    })


@require_POST
def cart_update(request, product_id):
    action = request.POST.get('action', 'inc')
    if action == 'inc' and not _get_store_config().is_open:
        return HttpResponseForbidden("Store is closed")

    cart = _get_cart(request)
    key = str(product_id)

    if key in cart:
        if action == 'inc':
            cart[key] = int(cart[key]) + 1
        elif action == 'dec':
            cart[key] = int(cart[key]) - 1
            if cart[key] <= 0:
                del cart[key]
        elif action == 'remove':
            del cart[key]

    _save_cart(request, cart)

    items, subtotal, delivery_fee, grand_total, amount_for_free_delivery, config = _cart_details(request)
    response = render(request, 'marketplace/_cart_body.html', {
        'items': items,
        'subtotal': subtotal,
        'delivery_fee': delivery_fee,
        'grand_total': grand_total,
        'amount_for_free_delivery': amount_for_free_delivery,
        'config': config,
    })
    response['HX-Trigger'] = 'cartUpdated'
    return response


@store_open_required
def checkout(request):
    apartment_id = request.POST.get('apartment') or request.GET.get('apartment')
    items, subtotal, delivery_fee, grand_total, amount_for_free_delivery, config = _cart_details(
        request,
        apartment_id,
    )
    if not items:
        messages.warning(request, 'Your cart is empty.')
        return redirect('home')
    if subtotal < config.minimum_order_value:
        messages.error(
            request,
            f'Minimum order value is UGX {config.minimum_order_value:,.0f}. Please add more items.',
        )
        return redirect('cart_view')

    apartments = Apartment.objects.all().order_by('name')
    saved_customer = None
    customer_id = request.session.get('customer_id')
    if customer_id:
        try:
            saved_customer = Customer.objects.get(id=customer_id)
        except (Customer.DoesNotExist, ValueError, TypeError):
            request.session.pop('customer_id', None)

    if request.method == 'POST':
        phone = request.POST.get('phone', '').strip()
        name = request.POST.get('name', '').strip()
        apartment_id = request.POST.get('apartment', '').strip()
        unit_number = request.POST.get('unit_number', '').strip()
        payment_method = request.POST.get('payment_method', '').strip()
        note = request.POST.get('note', '').strip()

        apartment = None
        if not phone:
            messages.error(request, 'Please enter your phone number.')
        if not apartment_id:
            messages.error(request, 'Please select your apartment.')
        else:
            try:
                apartment = Apartment.objects.get(pk=apartment_id)
            except (Apartment.DoesNotExist, ValueError, TypeError):
                messages.error(request, 'Please select a valid apartment.')
        if payment_method not in {'cash', 'momo'}:
            messages.error(request, 'Please select a valid payment method.')

        if not phone or apartment is None or payment_method not in {'cash', 'momo'}:
            return render(request, 'marketplace/checkout.html', {
                'items': items,
                'subtotal': subtotal,
                'delivery_fee': delivery_fee,
                'grand_total': grand_total,
                'amount_for_free_delivery': amount_for_free_delivery,
                'config': config,
                'apartments': apartments,
                'form_data': request.POST,
                'item_count': sum(item['quantity'] for item in items),
            })

        with transaction.atomic():
            customer, _ = Customer.objects.update_or_create(
                phone_number=phone,
                defaults={
                    'name': name,
                    'apartment': apartment,
                    'unit_number': unit_number,
                },
            )
            request.session['customer_id'] = customer.id
            order = Order.objects.create(
                customer=customer,
                status='new',
                payment_method=payment_method,
                total_amount=grand_total,
                delivery_fee=delivery_fee,
                note=note,
            )
            for item in items:
                OrderItem.objects.create(
                    order=order,
                    product=item['product'],
                    quantity=item['quantity'],
                    price_at_time_of_order=item['product'].selling_price,
                )

            request.session['cart'] = {}
            request.session.modified = True

        return redirect('order_success', order_id=order.id)

    form_data = {}
    if saved_customer:
        form_data = {
            'phone': saved_customer.phone_number,
            'name': saved_customer.name,
            'apartment': str(saved_customer.apartment.id) if saved_customer.apartment else '',
            'unit_number': saved_customer.unit_number,
        }
        if saved_customer.apartment:
            apartment_id = str(saved_customer.apartment_id)
            items, subtotal, delivery_fee, grand_total, amount_for_free_delivery, config = _cart_details(
                request,
                apartment_id,
            )

    return render(request, 'marketplace/checkout.html', {
        'items': items,
        'subtotal': subtotal,
        'delivery_fee': delivery_fee,
        'grand_total': grand_total,
        'amount_for_free_delivery': amount_for_free_delivery,
        'config': config,
        'apartments': apartments,
        'form_data': form_data,
        'item_count': sum(item['quantity'] for item in items),
    })


def checkout_summary_partial(request):
    apartment_id = request.GET.get('apartment')
    items, subtotal, delivery_fee, grand_total, amount_for_free_delivery, _config = _cart_details(
        request,
        apartment_id,
    )
    return render(request, 'marketplace/_checkout_summary.html', {
        'items': items,
        'subtotal': subtotal,
        'delivery_fee': delivery_fee,
        'grand_total': grand_total,
        'amount_for_free_delivery': amount_for_free_delivery,
    })


@login_required
def staff_settings(request):
    config = _get_store_config()

    if request.method == 'POST':
        try:
            amount_field = forms.DecimalField(
                max_digits=10,
                decimal_places=2,
                min_value=0,
            )
            values = {
                field: amount_field.clean(request.POST.get(field, ''))
                for field in (
                    'delivery_fee',
                    'free_delivery_threshold',
                    'minimum_order_value',
                )
            }
        except ValidationError:
            messages.error(
                request,
                'Enter valid non-negative amounts with no more than two decimal places.',
            )
        else:
            for field, value in values.items():
                setattr(config, field, value)
            config.closed_message = request.POST.get(
                'closed_message',
                config.closed_message,
            )[:255]
            config.whatsapp_number = re.sub(
                r'\D',
                '',
                request.POST.get('whatsapp_number', ''),
            )[:20]
            config.save()
            messages.success(request, 'Settings updated successfully.')
            return redirect('staff_settings')

    return render(request, 'marketplace/staff_settings.html', {'config': config})


@login_required
@require_POST
def staff_upload_stock(request):
    stock_file = request.FILES.get('stock_file')
    if stock_file is None:
        messages.error(request, 'No file uploaded.')
        return redirect('staff_settings')
    if not stock_file.name.lower().endswith('.xlsx'):
        messages.error(request, 'Please upload an Excel (.xlsx) file.')
        return redirect('staff_settings')

    try:
        workbook = openpyxl.load_workbook(
            stock_file,
            data_only=True,
            read_only=True,
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        logger.warning('Unable to open uploaded stock workbook: %s', exc)
        messages.error(request, 'Could not read the uploaded Excel file.')
        return redirect('staff_settings')

    try:
        try:
            worksheet = workbook['StockSummary']
        except KeyError:
            messages.error(
                request,
                "The uploaded file must contain a 'StockSummary' worksheet.",
            )
            return redirect('staff_settings')

        header_row = None
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=30, values_only=True),
            start=1,
        ):
            if row and str(row[0]).strip() == 'Product Code':
                header_row = row_number
                break

        if header_row is None:
            messages.error(
                request,
                "Could not find 'Product Code' header in the uploaded file.",
            )
            return redirect('staff_settings')

        stock_data = {}
        skipped_count = 0
        quantity_field = forms.DecimalField(max_digits=10, decimal_places=2)
        for row in worksheet.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ):
            if not row or len(row) <= 3:
                continue

            code = row[0]
            if code is None or not str(code).strip():
                continue

            product_code = str(code).strip()
            if product_code.startswith(('Product Group', 'Total')):
                continue

            raw_quantity = row[3]
            if isinstance(raw_quantity, str):
                raw_quantity = raw_quantity.replace(',', '').strip()
            if raw_quantity in (None, ''):
                skipped_count += 1
                continue

            try:
                quantity = quantity_field.clean(raw_quantity)
            except (ValidationError, InvalidOperation, TypeError, ValueError):
                skipped_count += 1
                continue
            stock_data[product_code] = quantity
    finally:
        workbook.close()

    updated_products = []
    with transaction.atomic():
        products = Product.objects.filter(code__in=stock_data)
        for product in products:
            new_quantity = stock_data[product.code]
            new_is_active = (
                False if new_quantity <= 0 else not product.manually_disabled
            )
            if (
                product.stock_quantity != new_quantity
                or product.is_active != new_is_active
            ):
                product.stock_quantity = new_quantity
                product.is_active = new_is_active
                updated_products.append(product)

        if updated_products:
            Product.objects.bulk_update(
                updated_products,
                ['stock_quantity', 'is_active'],
                batch_size=500,
            )

    messages.success(
        request,
        f'Successfully updated stock for {len(updated_products)} products.',
    )
    if skipped_count:
        messages.warning(
            request,
            f'Skipped {skipped_count} stock row(s) with invalid quantities.',
        )
    return redirect('staff_settings')


def order_success(request, order_id):
    order = get_object_or_404(Order, id=order_id)
    return render(request, 'marketplace/order_success.html', {
        'order': order,
        'config': _get_store_config(),
        'whatsapp_message': f'Hi GGM Store, I need help with order #{order.id}',
    })


def customer_orders(request):
    customer_id = request.session.get('customer_id')
    if not customer_id:
        messages.info(request, 'Please place an order to see your order history.')
        return redirect('home')

    try:
        customer = Customer.objects.get(id=customer_id)
    except (Customer.DoesNotExist, ValueError, TypeError):
        request.session.pop('customer_id', None)
        messages.info(request, 'Please place an order to see your order history.')
        return redirect('home')

    orders = Order.objects.filter(customer=customer).order_by('-created_at').prefetch_related('items')
    return render(request, 'marketplace/customer_orders.html', {
        'customer': customer,
        'orders': orders,
    })


def customer_order_detail(request, order_id):
    customer_id = request.session.get('customer_id')
    if not customer_id:
        return redirect('home')

    try:
        order = get_object_or_404(
            Order.objects.prefetch_related('items__product'),
            id=order_id,
            customer_id=customer_id,
        )
    except (ValueError, TypeError):
        request.session.pop('customer_id', None)
        return redirect('home')

    if request.headers.get('HX-Request') == 'true':
        return render(
            request,
            'marketplace/_order_status_partial.html',
            {'order': order},
        )

    return render(request, 'marketplace/customer_order_detail.html', {
        'order': order,
        'config': _get_store_config(),
        'whatsapp_message': f'Hi GGM Store, I need help with order #{order.id}',
    })


@login_required
def staff_orders(request):
    status = request.GET.get('status', 'new')
    valid_statuses = {'new', 'confirmed', 'packed', 'delivered', 'all'}
    if status not in valid_statuses:
        status = 'new'

    orders = Order.objects.select_related(
        'customer', 'customer__apartment'
    ).prefetch_related('items__product')
    if status != 'all':
        orders = orders.filter(status=status)
    orders = orders.order_by('-created_at')

    context = {
        'status': status,
        'orders': orders,
        'new_count': Order.objects.filter(status='new').count(),
        'confirmed_count': Order.objects.filter(status='confirmed').count(),
        'packed_count': Order.objects.filter(status='packed').count(),
        'delivered_count': Order.objects.filter(status='delivered').count(),
        'config': _get_store_config(),
    }
    template = (
        'marketplace/_staff_order_list.html'
        if request.headers.get('HX-Request') == 'true'
        else 'marketplace/staff_orders.html'
    )
    response = render(request, template, context)
    if request.headers.get('HX-Request') == 'true':
        response['HX-Trigger-After-Swap'] = json.dumps({
            'staffOrdersUpdated': {'newCount': context['new_count']},
        })
    return response


@staff_member_required
@require_POST
def staff_toggle_store_status(request):
    config = _get_store_config()
    config.is_open = not config.is_open
    config.save(update_fields=['is_open'])
    return render(
        request,
        'marketplace/_store_status_banner.html',
        {'config': config},
    )


@login_required
def staff_order_detail(request, order_id):
    order = get_object_or_404(Order, id=order_id)
    status_to_notify = {
        'new': 'confirmed',
        'packed': 'packed',
    }.get(order.status)
    customer_name = order.customer.name or 'there'
    whatsapp_message = ''
    if status_to_notify == 'confirmed':
        whatsapp_message = (
            f'Hi {customer_name}, your order #{order.id} at GGM Store has '
            f'been confirmed and is being packed! Total to pay: '
            f'UGX {order.total_amount:,.0f}'
        )
    elif status_to_notify == 'packed':
        whatsapp_message = (
            f'Hi {customer_name}, your order #{order.id} from GGM Store is '
            f'on the way! The driver is carrying it. Total to pay: '
            f'UGX {order.total_amount:,.0f}'
        )
    customer_whatsapp_number = re.sub(r'\D', '', order.customer.phone_number)
    if customer_whatsapp_number.startswith('0'):
        customer_whatsapp_number = f'256{customer_whatsapp_number[1:]}'
    return render(request, 'marketplace/staff_order_detail.html', {
        'order': order,
        'config': _get_store_config(),
        'status_to_notify': status_to_notify,
        'whatsapp_message': whatsapp_message,
        'customer_whatsapp_number': customer_whatsapp_number,
    })


@staff_member_required
@require_POST
def log_whatsapp_notification(request, order_id, status_sent):
    if status_sent not in {'confirmed', 'packed'}:
        return HttpResponse(status=400)

    order = get_object_or_404(Order, id=order_id)
    if status_sent == 'confirmed' and order.status != 'new':
        return HttpResponse(status=409)
    if status_sent == 'packed' and order.status != 'packed':
        return HttpResponse(status=409)

    WhatsAppNotificationLog.objects.create(
        order=order,
        status_sent=status_sent,
    )
    return HttpResponse(status=204)


@login_required
@require_POST
def staff_order_action(request, order_id):
    order = get_object_or_404(Order, id=order_id)
    action_statuses = {
        'confirm': 'confirmed',
        'packed': 'packed',
        'delivered': 'delivered',
        'cancel': 'cancelled',
    }
    action = request.POST.get('action')
    if action not in action_statuses:
        messages.error(request, 'Please select a valid order action.')
        return redirect('staff_order_detail', order_id=order.id)

    update_fields = ['status', 'updated_at']
    order.status = action_statuses[action]
    if action == 'confirm':
        payment_method = request.POST.get('payment_method')
        if payment_method:
            if payment_method not in {'cash', 'momo'}:
                messages.error(request, 'Please select a valid payment method.')
                return redirect('staff_order_detail', order_id=order.id)
            order.payment_method = payment_method
            update_fields.append('payment_method')
    order.save(update_fields=update_fields)

    if action == 'confirm':
        _send_push_to_customer(
            customer=order.customer,
            title='Order Confirmed!',
            body=(
                f"Hi {order.customer.name or 'there'}, your order #{order.id} "
                'has been confirmed and is being packed.'
            ),
            url=f'/my-orders/{order.id}/',
        )
    elif action == 'packed':
        _send_push_to_customer(
            customer=order.customer,
            title='Order is on the way! 🚚',
            body=(
                f'Your order #{order.id} is out for delivery. Please have '
                f'UGX {order.total_amount:,.0f} ready.'
            ),
            url=f'/my-orders/{order.id}/',
        )

    return redirect('staff_order_detail', order_id=order.id)


@login_required
def staff_stock(request):
    q = request.GET.get('q', '').strip()
    filter_type = request.GET.get('filter', '')
    products = Product.objects.all().order_by('name')
    if q:
        products = products.filter(
            Q(name__icontains=q) | Q(code__icontains=q)
        )

    if filter_type == 'low_stock':
        products = products.filter(
            is_active=True,
            stock_quantity__lte=F('low_stock_threshold'),
        ).order_by('stock_quantity', 'name')
    elif q:
        products = products[:50]

    return render(request, 'marketplace/staff_stock.html', {
        'products': products,
        'q': q,
        'filter_type': filter_type,
    })


@login_required
@require_POST
def staff_toggle_stock(request, product_id):
    product = get_object_or_404(Product, id=product_id)
    if product.manually_disabled:
        product.manually_disabled = False
        product.is_active = product.stock_quantity > 0
    else:
        product.manually_disabled = True
        product.is_active = False
    product.save(update_fields=['manually_disabled', 'is_active'])
    return render(request, 'marketplace/_staff_product_row.html', {'product': product})


@login_required
def staff_order_slip(request, order_id):
    order = get_object_or_404(
        Order.objects.select_related(
            'customer', 'customer__apartment'
        ).prefetch_related('items__product'),
        id=order_id,
    )
    return render(request, 'marketplace/staff_order_slip.html', {'order': order})


def service_worker(request):
    sw_path = settings.BASE_DIR / 'static' / 'sw.js'
    if not sw_path.exists():
        raise Http404
    response = FileResponse(sw_path.open('rb'), content_type='application/javascript')
    response['Service-Worker-Allowed'] = '/'
    response['Cache-Control'] = 'no-cache'
    return response


@require_POST
def save_push_subscription(request):
    if request.content_type != 'application/json':
        return JsonResponse(
            {'status': 'error', 'message': 'Expected application/json.'},
            status=415,
        )
    if len(request.body) > 4096:
        return JsonResponse(
            {'status': 'error', 'message': 'Subscription payload is too large.'},
            status=400,
        )

    try:
        data = json.loads(request.body)
    except (json.JSONDecodeError, UnicodeDecodeError):
        return JsonResponse(
            {'status': 'error', 'message': 'Invalid JSON.'},
            status=400,
        )
    if not isinstance(data, dict):
        return JsonResponse(
            {'status': 'error', 'message': 'Invalid subscription data.'},
            status=400,
        )

    endpoint = data.get('endpoint')
    keys = data.get('keys')
    if not isinstance(endpoint, str) or not isinstance(keys, dict):
        return JsonResponse(
            {'status': 'error', 'message': 'Missing subscription keys.'},
            status=400,
        )

    p256dh = keys.get('p256dh')
    auth = keys.get('auth')
    if (
        not endpoint
        or len(endpoint) > 500
        or not isinstance(p256dh, str)
        or not 1 <= len(p256dh) <= 200
        or not isinstance(auth, str)
        or not 1 <= len(auth) <= 100
    ):
        return JsonResponse(
            {'status': 'error', 'message': 'Missing or invalid subscription keys.'},
            status=400,
        )

    try:
        URLValidator(schemes=['https'])(endpoint)
    except ValidationError:
        return JsonResponse(
            {'status': 'error', 'message': 'Invalid subscription endpoint.'},
            status=400,
        )

    customer_id = request.session.get('customer_id')
    customer = (
        Customer.objects.filter(id=customer_id).first()
        if customer_id
        else None
    )
    PushSubscription.objects.update_or_create(
        endpoint=endpoint,
        defaults={
            'p256dh': p256dh,
            'auth': auth,
            'customer': customer,
        },
    )
    return JsonResponse({'status': 'success'})


@staff_member_required
def owner_dashboard(request):
    today = timezone.localdate()
    date_str = request.GET.get('date')
    try:
        selected_date = datetime.strptime(date_str, '%Y-%m-%d').date() if date_str else today
    except ValueError:
        selected_date = today

    day_orders = Order.objects.filter(
        created_at__date=selected_date,
    ).exclude(status='cancelled')

    total_orders = day_orders.count()
    total_revenue = day_orders.aggregate(total=Sum('total_amount'))['total'] or 0
    avg_order_value = total_revenue / total_orders if total_orders else 0

    cash_orders = day_orders.filter(payment_method='cash').count()
    momo_orders = day_orders.filter(payment_method='momo').count()
    top_products = OrderItem.objects.filter(
        order__in=day_orders,
        product__isnull=False,
    ).values(
        'product__name',
    ).annotate(
        total_qty=Sum('quantity'),
    ).order_by('-total_qty', 'product__name')[:5]

    seven_days_ago = selected_date - timedelta(days=6)
    trend_data = Order.objects.filter(
        created_at__date__gte=seven_days_ago,
        created_at__date__lte=selected_date,
    ).exclude(
        status='cancelled',
    ).values(
        'created_at__date',
    ).annotate(
        daily_revenue=Sum('total_amount'),
    ).order_by('created_at__date')
    trend_revenue_by_date = {
        item['created_at__date']: item['daily_revenue']
        for item in trend_data
    }
    seven_day_trend = [
        {
            'date': day,
            'revenue': trend_revenue_by_date.get(day, 0),
            'is_selected': day == selected_date,
        }
        for day in (
            seven_days_ago + timedelta(days=offset)
            for offset in range(7)
        )
    ]
    max_trend_revenue = max(
        (day['revenue'] for day in seven_day_trend),
        default=0,
    ) or 1

    recent_orders = day_orders.select_related(
        'customer',
    ).order_by('-created_at')[:10]

    low_stock_products = Product.objects.filter(
        is_active=True,
        stock_quantity__lte=F('low_stock_threshold'),
    ).order_by('stock_quantity', 'name')[:5]

    return render(request, 'marketplace/owner_dashboard.html', {
        'selected_date': selected_date,
        'today_date': today,
        'yesterday_date': today - timedelta(days=1),
        'is_today': selected_date == today,
        'total_orders': total_orders,
        'total_revenue': total_revenue,
        'avg_order_value': avg_order_value,
        'cash_orders': cash_orders,
        'momo_orders': momo_orders,
        'top_products': top_products,
        'low_stock_products': low_stock_products,
        'recent_orders': recent_orders,
        'seven_day_trend': seven_day_trend,
        'max_trend_revenue': max_trend_revenue,
    })


@login_required
def staff_manifest(request):
    today = timezone.localdate()
    orders = Order.objects.filter(
        created_at__date=today,
        status__in=['confirmed', 'packed'],
    ).select_related(
        'customer',
        'customer__apartment',
    ).order_by(
        'customer__apartment__name',
        'id',
    )
    totals = orders.aggregate(
        cash_to_collect=Sum('total_amount', filter=Q(payment_method='cash')),
        momo_to_collect=Sum('total_amount', filter=Q(payment_method='momo')),
    )

    return render(request, 'marketplace/staff_manifest.html', {
        'orders': orders,
        'today': today,
        'cash_to_collect': totals['cash_to_collect'] or 0,
        'momo_to_collect': totals['momo_to_collect'] or 0,
        'total_orders': orders.count(),
    })