import openpyxl
import os
from django.core.management.base import BaseCommand
from tqdm import tqdm
from marketplace.models import ProductGroup, Unit, Product


def to_number(value):
    """Convert Excel cell value to float, handling commas, strings, and None."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    # It's a string — clean it
    cleaned = str(value).strip().replace(',', '')
    if cleaned == '' or cleaned.lower() == 'null':
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


class Command(BaseCommand):
    help = 'Import products and stock from Excel files'

    def add_arguments(self, parser):
        parser.add_argument('--skip-products', action='store_true', help='Skip product import')
        parser.add_argument('--skip-stock', action='store_true', help='Skip stock import')

    def handle(self, *args, **kwargs):
        product_file = 'data/products.xlsx'
        stock_file = 'data/stock.xlsx'

        if not kwargs['skip_products']:
            if not os.path.exists(product_file):
                self.stdout.write(self.style.ERROR(f'File not found: {product_file}'))
            else:
                self.stdout.write(self.style.WARNING('Starting product import...'))
                self.import_products(product_file)

        if not kwargs['skip_stock']:
            if not os.path.exists(stock_file):
                self.stdout.write(self.style.ERROR(f'File not found: {stock_file}'))
            else:
                self.stdout.write(self.style.WARNING('Starting stock import...'))
                self.import_stock(stock_file)

        self.stdout.write(self.style.SUCCESS('Done.'))

    def import_products(self, filepath):
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        ws = wb['Product']

        total_rows = ws.max_row or 0
        # Data starts at row 2 (row 1 = headers)
        data_rows = max(total_rows - 1, 0)

        created_count = 0
        updated_count = 0
        skipped_count = 0

        with tqdm(total=data_rows, desc='Products', unit='row', ncols=80) as pbar:
            for row in ws.iter_rows(min_row=2, values_only=True):
                code = row[0]
                name = row[1]
                group_name = row[2]
                unit_name = row[5]
                cost_price = row[7] if row[7] is not None else 0
                selling_price = row[8] if row[8] is not None else 0

                if not code or not name:
                    skipped_count += 1
                    pbar.update(1)
                    continue

                group = None
                if group_name and str(group_name).strip() and str(group_name).lower() != 'null':
                    group, _ = ProductGroup.objects.get_or_create(name=str(group_name).strip())

                unit = None
                if unit_name and str(unit_name).strip():
                    unit, _ = Unit.objects.get_or_create(name=str(unit_name).strip())

                manually_disabled = False
                if not group_name or str(group_name).lower() == 'null':
                    manually_disabled = True
                if 'disabled' in str(name).lower() or 'invalid' in str(name).lower():
                    manually_disabled = True
                try:
                    if float(selling_price) <= 0:
                        manually_disabled = True
                except (TypeError, ValueError):
                    manually_disabled = True

                _, created = Product.objects.update_or_create(
                    code=str(code).strip(),
                    defaults={
                        'name': str(name).strip(),
                        'group': group,
                        'unit': unit,
                        'cost_price': cost_price or 0,
                        'selling_price': selling_price or 0,
                        'manually_disabled': manually_disabled,
                    }
                )
                if created:
                    created_count += 1
                else:
                    updated_count += 1

                pbar.update(1)
                pbar.set_postfix(created=created_count, updated=updated_count, skipped=skipped_count)

        wb.close()
        self.stdout.write(self.style.SUCCESS(
            f'Products done. Created: {created_count}, Updated: {updated_count}, Skipped: {skipped_count}'
        ))

    def import_stock(self, filepath):
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        ws = wb['StockSummary']

        # Locate the header row where "Product Code" is in column A
        header_row = None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
            if row and row[0] == 'Product Code':
                header_row = i
                break

        if not header_row:
            self.stdout.write(self.style.ERROR('Could not find "Product Code" header in stock file.'))
            wb.close()
            return

        total_rows = ws.max_row or 0
        data_rows = max(total_rows - header_row, 0)

        updated_count = 0
        missing_count = 0
        skipped_bad_qty = 0

        with tqdm(total=data_rows, desc='Stock   ', unit='row', ncols=80) as pbar:
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                code = row[0]
                quantity_raw = row[3] if len(row) > 3 else None

                # Skip blank rows
                if not code or str(code).strip() == '':
                    pbar.update(1)
                    continue

                code_str = str(code).strip()
                # Skip group headers and total rows
                if code_str.startswith('Product Group') or code_str.startswith('Total'):
                    pbar.update(1)
                    continue

                quantity = to_number(quantity_raw)
                if quantity is None:
                    skipped_bad_qty += 1
                    pbar.update(1)
                    continue

                try:
                    product = Product.objects.get(code=code_str)
                except Product.DoesNotExist:
                    missing_count += 1
                    pbar.update(1)
                    continue

                product.stock_quantity = quantity
                if quantity <= 0:
                    product.is_active = False
                elif not product.manually_disabled:
                    product.is_active = True

                product.save(update_fields=['stock_quantity', 'is_active'])
                updated_count += 1
                pbar.update(1)
                pbar.set_postfix(updated=updated_count, missing=missing_count, bad_qty=skipped_bad_qty)

        wb.close()
        self.stdout.write(self.style.SUCCESS(
            f'Stock done. Updated: {updated_count}, '
            f'Skipped (not in catalog): {missing_count}, '
            f'Skipped (bad quantity): {skipped_bad_qty}'
        ))